"""File extraction: PDF / Word / Excel -> clean, LINE-PRESERVED text.

Why line-preserved and not one big string? The single most important lesson in
this project: run NER **per line**, not on the whole document. Tables with numbers
between rows wreck recall if you feed the model one long noisy sequence. So the
extractor's job is to hand downstream code a list of lines, each tagged with where
it came from (page / paragraph / sheet cell), while still being able to render a
plain-text view identical to what detection runs on.

Public API:
    extract(path)      -> ExtractedDoc   (dispatch by file extension)
    ExtractedDoc.text  -> str            (lines joined by newline; what the UI shows
                                           and what detection runs on)
    ExtractedDoc.lines -> list[Line]     (each with .text and .source)

PDF word geometry for in-place redaction is intentionally NOT here — that lives in
redact_pdf.py using PyMuPDF, because redaction needs (block_no, line_no) coordinates
that pdfplumber's text view discards.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Line:
    """One line of source text and where it came from."""
    text: str
    source: str  # e.g. "p1" (PDF page 1), "para12", "table1:r2c3", "Sheet1!B4"


@dataclass
class ExtractedDoc:
    """A document reduced to line-preserved text, ready for per-line analysis."""
    lines: list[Line] = field(default_factory=list)
    filetype: str = ""
    path: str = ""

    @property
    def text(self) -> str:
        """The full document as newline-joined lines. This is the canonical text
        that detection runs on and that the UI displays — they must match exactly."""
        return "\n".join(line.text for line in self.lines)

    def __len__(self) -> int:
        return len(self.lines)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def extract_pdf(path: str | Path) -> ExtractedDoc:
    """Extract text from a PDF with pdfplumber, one Line per visual text line.

    We keep pdfplumber's line grouping (its layout heuristics already split rows),
    which preserves table rows as separate lines — exactly what per-line NER needs.
    """
    import pdfplumber

    path = Path(path)
    doc = ExtractedDoc(filetype="pdf", path=str(path))
    with pdfplumber.open(str(path)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            for raw in text.split("\n"):
                stripped = raw.rstrip()
                if stripped.strip():
                    doc.lines.append(Line(text=stripped, source=f"p{page_no}"))
    return doc


# --------------------------------------------------------------------------- #
# Word (.docx)
# --------------------------------------------------------------------------- #
def extract_docx(path: str | Path) -> ExtractedDoc:
    """Extract text from a Word document: body paragraphs AND table cells.

    Tables are the whole point of this spike, so we walk them explicitly. Each
    non-empty paragraph and each non-empty table cell becomes its own Line, tagged
    with a source that later lets us write the anonymized value back to the right
    place (see anonymize_files.py).
    """
    from docx import Document
    from docx.document import Document as _DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    path = Path(path)
    doc = ExtractedDoc(filetype="docx", path=str(path))
    document = Document(str(path))

    # Walk body in document order so paragraphs and tables interleave correctly.
    para_idx = 0
    table_idx = 0
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            text = para.text.strip()
            if text:
                doc.lines.append(Line(text=text, source=f"para{para_idx}"))
            para_idx += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    text = cell.text.strip()
                    if text:
                        doc.lines.append(
                            Line(text=text, source=f"table{table_idx}:r{r}c{c}")
                        )
            table_idx += 1
    return doc


# --------------------------------------------------------------------------- #
# Excel (.xlsx)
# --------------------------------------------------------------------------- #
def extract_xlsx(path: str | Path) -> ExtractedDoc:
    """Extract text from an Excel workbook: one Line per non-empty cell.

    Cell-per-line mirrors the per-line NER lesson for spreadsheets. The source tag
    (e.g. 'Sheet1!B4') lets anonymize_files.py write the placeholder back to the
    exact cell. Numbers are stringified so numeric PII (SSN typed as a number,
    account numbers) is still seen by the detectors.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    path = Path(path)
    doc = ExtractedDoc(filetype="xlsx", path=str(path))
    wb = load_workbook(str(path), read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if not text:
                    continue
                col = get_column_letter(cell.column)
                source = f"{ws.title}!{col}{cell.row}"
                doc.lines.append(Line(text=text, source=source))
    wb.close()
    return doc


# --------------------------------------------------------------------------- #
# Dispatch
# --------------------------------------------------------------------------- #
_EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
}


def extract(path: str | Path) -> ExtractedDoc:
    """Extract a supported file to line-preserved text, chosen by extension."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in _EXTRACTORS:
        raise ValueError(
            f"Unsupported file type {ext!r}. Supported: {sorted(_EXTRACTORS)}"
        )
    if not path.exists():
        raise FileNotFoundError(f"No such file: {path}")
    return _EXTRACTORS[ext](path)


SUPPORTED_EXTENSIONS = tuple(_EXTRACTORS)
