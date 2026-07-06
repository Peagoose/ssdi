"""Generate MINIMAL synthetic fixtures (fake PII) for exercising extract.py.

These are tiny, throwaway files used by tests to prove extraction works across
PDF / DOCX / XLSX, including a table in each. The full labeled test set for recall
measurement is built later (Task 9). All PII here is fake and non-sequential.

Usage:  python -m tests.make_min_fixtures <out_dir>
"""
from __future__ import annotations

import sys
from pathlib import Path


def make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("IN THE UNITED STATES DISTRICT COURT")
    doc.add_paragraph("Plaintiff Sarah Chen brings this action.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "SSN"
    table.cell(1, 0).text = "Robert Alvarez"
    table.cell(1, 1).text = "536-90-4788"
    doc.save(str(path))


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Parties"
    ws.append(["Name", "SSN", "Phone"])
    ws.append(["Robert Alvarez", "536-90-4788", "(212) 555-0182"])
    ws.append(["Maria Gomez", "457-55-1462", "(718) 555-0147"])
    wb.save(str(path))


def make_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    text = c.beginText(72, 720)
    for line in [
        "IN THE UNITED STATES DISTRICT COURT",
        "Plaintiff Sarah Chen brings this action.",
        "Name            SSN            Phone",
        "Robert Alvarez  536-90-4788    (212) 555-0182",
    ]:
        text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("test_files")
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "min.docx")
    make_xlsx(out / "min.xlsx")
    make_pdf(out / "min.pdf")
    print(f"Wrote min.docx, min.xlsx, min.pdf to {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
