"""Tests for redact_pdf.py and anonymize_files.py — real files, no NER model.

We construct entities/anonymization directly so these run without en_core_web_lg,
and verify the OUTPUT files: originals removed, placeholders present.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from src.analyze import Entity
from src.anonymize import anonymize
from src.extract import extract
from src.anonymize_files import anonymize_docx, anonymize_xlsx
from src.redact_pdf import redact_values, _find_runs, _Word
from tests.make_min_fixtures import make_docx, make_pdf, make_xlsx


# --------------------------------------------------------------------------- #
# Contiguous-run matcher (the core PDF-redaction trap)
# --------------------------------------------------------------------------- #
def _w(text, block=0, line=0, x0=0.0):
    return _Word(x0, 0.0, x0 + 10, 10.0, text, block, line)


def test_find_runs_exact_number_not_substring():
    # "4788" must NOT match inside "536-90-4788"
    words = [_w("536-90-4788", x0=0)]
    assert list(_find_runs(words, "4788")) == []
    assert list(_find_runs(words, "536-90-4788")) == [(0, 1)]


def test_find_runs_multiword_name():
    words = [_w("Robert", x0=0), _w("Alvarez", x0=20), _w("signed", x0=40)]
    assert list(_find_runs(words, "Robert Alvarez")) == [(0, 2)]


def test_find_runs_number_split_across_words():
    # digits spread across a run concatenate exactly
    words = [_w("536", x0=0), _w("90", x0=20), _w("4788", x0=40)]
    assert list(_find_runs(words, "536-90-4788")) == [(0, 3)]


# --------------------------------------------------------------------------- #
# Real PDF redaction: value is truly removed
# --------------------------------------------------------------------------- #
def test_pdf_redaction_removes_value(tmp_path: Path):
    fitz = pytest.importorskip("fitz")
    src = tmp_path / "in.pdf"
    make_pdf(src)
    out = tmp_path / "out.pdf"

    count = redact_values(src, out, [("536-90-4788", "[SSN_1]")])
    assert count >= 1

    doc = fitz.open(str(out))
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "536-90-4788" not in text          # original truly gone
    assert "[SSN_1]" in text                   # placeholder stamped in


def test_pdf_redaction_leaves_other_text(tmp_path: Path):
    fitz = pytest.importorskip("fitz")
    src = tmp_path / "in.pdf"
    make_pdf(src)
    out = tmp_path / "out.pdf"
    redact_values(src, out, [("536-90-4788", "[SSN_1]")])
    doc = fitz.open(str(out))
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "Robert Alvarez" in text            # untargeted text preserved


# --------------------------------------------------------------------------- #
# DOCX / XLSX write-back
# --------------------------------------------------------------------------- #
def test_docx_write_back(tmp_path: Path):
    src = tmp_path / "in.docx"
    make_docx(src)
    doc = extract(src)
    # Build entities for the two names + the SSN from the extracted lines
    ents = []
    for line in doc.lines:
        pass
    text = doc.text
    for etype, value in [("PERSON", "Sarah Chen"), ("PERSON", "Robert Alvarez"),
                         ("US_SSN", "536-90-4788")]:
        idx = text.index(value)
        ents.append(Entity(etype, idx, idx + len(value), 1.0, value,
                           source=doc.lines[text.count("\n", 0, idx)].source))
    anon = anonymize(text, ents)

    out = tmp_path / "out.docx"
    anonymize_docx(src, out, doc, anon)

    from docx import Document
    written = Document(str(out))
    body = "\n".join(p.text for p in written.paragraphs)
    for tbl in written.tables:
        for row in tbl.rows:
            for cell in row.cells:
                body += "\n" + cell.text
    assert "536-90-4788" not in body
    assert "Sarah Chen" not in body
    assert "[SSN_1]" in body or "[NAME_" in body


def test_xlsx_write_back(tmp_path: Path):
    src = tmp_path / "in.xlsx"
    make_xlsx(src)
    doc = extract(src)
    text = doc.text
    ents = []
    for etype, value in [("US_SSN", "536-90-4788"), ("PERSON", "Robert Alvarez")]:
        idx = text.index(value)
        ents.append(Entity(etype, idx, idx + len(value), 1.0, value,
                           source=doc.lines[text.count("\n", 0, idx)].source))
    anon = anonymize(text, ents)

    out = tmp_path / "out.xlsx"
    anonymize_xlsx(src, out, doc, anon)

    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    values = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    wb.close()
    joined = " ".join(values)
    assert "536-90-4788" not in joined
    assert "[SSN_1]" in joined
