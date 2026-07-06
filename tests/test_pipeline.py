"""Tests for src/pipeline.py glue (no NER model needed)."""
from __future__ import annotations

from pathlib import Path

import pytest

from src.analyze import Entity
from src.anonymize import anonymize
from src.extract import extract
from src.pipeline import Processed, redaction_pairs, write_anonymized_output
from tests.make_min_fixtures import make_docx, make_pdf, make_xlsx


def _processed(path):
    doc = extract(path)
    text = doc.text
    ents = []
    for etype, value in [("US_SSN", "536-90-4788"), ("PERSON", "Robert Alvarez")]:
        if value in text:
            idx = text.index(value)
            ents.append(Entity(etype, idx, idx + len(value), 1.0, value,
                               source=doc.lines[text.count("\n", 0, idx)].source))
    anon = anonymize(text, ents)
    return Processed(doc=doc, entities=ents, anonymization=anon)


def test_redaction_pairs_unique():
    text = "536-90-4788 and 536-90-4788 again"
    ents = [Entity("US_SSN", 0, 11, 1.0, "536-90-4788"),
            Entity("US_SSN", 16, 27, 1.0, "536-90-4788")]
    anon = anonymize(text, ents)
    pairs = redaction_pairs(anon)
    assert pairs == [("536-90-4788", "[SSN_1]")]


def test_write_output_pdf(tmp_path: Path):
    pytest.importorskip("fitz")
    src = tmp_path / "in.pdf"
    make_pdf(src)
    processed = _processed(src)
    out = tmp_path / "out.pdf"
    write_anonymized_output(src, out, processed)
    import fitz
    d = fitz.open(str(out))
    text = "\n".join(p.get_text() for p in d)
    d.close()
    assert "536-90-4788" not in text


def test_write_output_docx(tmp_path: Path):
    src = tmp_path / "in.docx"
    make_docx(src)
    processed = _processed(src)
    out = tmp_path / "out.docx"
    write_anonymized_output(src, out, processed)
    from docx import Document
    doc = Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                body += "\n" + cell.text
    assert "536-90-4788" not in body


def test_write_output_xlsx(tmp_path: Path):
    src = tmp_path / "in.xlsx"
    make_xlsx(src)
    processed = _processed(src)
    out = tmp_path / "out.xlsx"
    write_anonymized_output(src, out, processed)
    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    joined = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    wb.close()
    assert "536-90-4788" not in joined
